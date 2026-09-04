"""文档解析模块。"""

from pathlib import Path
from typing import Optional


def parse_file(file_path: str, file_type: str = "") -> str:
    """解析文件并提取文本内容。

    Args:
        file_path: 文件路径
        file_type: 文件类型（扩展名），为空则自动推断

    Returns:
        提取的文本内容
    """
    path = Path(file_path)
    if not path.exists():
        return ""

    if not file_type:
        file_type = path.suffix.lstrip(".").lower()

    if file_type in ("txt", "md", "markdown"):
        return _parse_text(path)
    elif file_type == "pdf":
        return _parse_pdf(path)
    elif file_type == "docx":
        return _parse_docx(path)
    elif file_type == "json":
        return _parse_json(path)
    elif file_type in ("csv",):
        return _parse_csv(path)
    elif file_type in ("py", "js", "ts", "java", "go", "rs", "sh"):
        return _parse_text(path)
    else:
        # 未知格式，尝试以文本方式读取
        return _parse_text(path)


def _parse_text(path: Path) -> str:
    """解析纯文本文件。"""
    try:
        return path.read_text(encoding="utf-8")
    except UnicodeDecodeError:
        try:
            return path.read_text(encoding="gbk")
        except Exception:
            return ""


def _parse_json(path: Path) -> str:
    """解析 JSON 文件，返回格式化文本。"""
    import json

    try:
        data = json.loads(path.read_text(encoding="utf-8"))
        return json.dumps(data, ensure_ascii=False, indent=2)
    except Exception:
        return _parse_text(path)


def _parse_csv(path: Path) -> str:
    """解析 CSV 文件，返回文本。"""
    return _parse_text(path)


def _parse_pdf(path: Path) -> str:
    """提取 PDF 中每一页的文本。"""
    from pypdf import PdfReader

    try:
        pages = (page.extract_text() or "" for page in PdfReader(path).pages)
        return "\n\n".join(text.strip() for text in pages if text.strip())
    except Exception:
        return ""


def _parse_docx(path: Path) -> str:
    """提取 DOCX 的段落与表格文本。"""
    from docx import Document

    try:
        document = Document(path)
        blocks = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    blocks.append("\t".join(cells))
        return "\n".join(blocks)
    except Exception:
        return ""


def detect_file_type(filename: str) -> str:
    """从文件名推断文件类型。"""
    ext = Path(filename).suffix.lstrip(".").lower()
    return ext if ext else "unknown"


def get_file_size_str(size_bytes: int) -> str:
    """将字节大小转为可读字符串。"""
    if size_bytes < 1024:
        return f"{size_bytes} B"
    elif size_bytes < 1024 * 1024:
        return f"{size_bytes / 1024:.1f} KB"
    else:
        return f"{size_bytes / (1024 * 1024):.1f} MB"
